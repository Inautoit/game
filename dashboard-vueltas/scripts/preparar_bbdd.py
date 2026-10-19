"""Prepara las BBDD mensuales exportadas (Exportadas_<Mes><AA>*.xlsx).

- Añade la columna MES_BBDD (AAAA-MM) en cada pestaña del Excel.
- Genera data/bbdd_<AAAA-MM>.json con {mes, campanas: {<pestaña>: [instalaciones]}}
  (solo números de instalación, sin datos personales) para el dashboard.

Uso: python3 scripts/preparar_bbdd.py <salida_xlsx> <salida_json> fichero1.xlsx ...
"""
import json, os, re, sys
import openpyxl

MESES = {'enero': 1, 'febrero': 2, 'marzo': 3, 'abril': 4, 'mayo': 5, 'junio': 6, 'julio': 7,
         'agosto': 8, 'septiembre': 9, 'setiembre': 9, 'octubre': 10, 'noviembre': 11, 'diciembre': 12}


def mes_de_nombre(nombre):
    m = re.search(r'(' + '|'.join(MESES) + r')[ _-]?(\d{2,4})', nombre, re.I)
    if not m:
        raise SystemExit(f'No encuentro el mes en el nombre: {nombre}')
    y = int(m.group(2))
    return f'{2000 + y if y < 100 else y}-{MESES[m.group(1).lower()]:02d}'


def norm_id(v):
    s = str(v).strip()
    return s[:-2] if s.endswith('.0') else s


def main(out_xlsx, out_json, files):
    for f in files:
        mes = mes_de_nombre(os.path.basename(f))
        wb = openpyxl.load_workbook(f)
        campanas = {}
        for ws in wb.worksheets:
            hdr = [c.value for c in ws[1]]
            col_id = hdr.index('Numero_de_Instalacion') + 1
            ncol = len(hdr)
            while ncol and hdr[ncol - 1] is None:
                ncol -= 1
            col_mes = ncol + 1 if 'MES_BBDD' not in hdr else hdr.index('MES_BBDD') + 1
            ws.cell(row=1, column=col_mes, value='MES_BBDD')
            ids = []
            for r in range(2, ws.max_row + 1):
                v = ws.cell(row=r, column=col_id).value
                if v is None or str(v).strip() == '':
                    continue
                ws.cell(row=r, column=col_mes, value=mes)
                ids.append(norm_id(v))
            campanas[ws.title.strip()] = sorted(set(ids))
            print(f'{mes} {ws.title.strip():25s} filas={len(ids):6d} unicas={len(set(ids)):6d}')
        nombre = re.sub(r'^[0-9a-f]{8}-', '', os.path.basename(f))
        wb.save(os.path.join(out_xlsx, nombre))
        with open(os.path.join(out_json, f'bbdd_{mes}.json'), 'w') as fh:
            json.dump({'mes': mes, 'campanas': campanas}, fh, separators=(',', ':'))


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2], sys.argv[3:])
